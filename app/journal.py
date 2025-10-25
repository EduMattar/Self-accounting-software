from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .accounts import AccountCatalog
from .ledger import LedgerEntry


@dataclass
class JournalRow:
    transaction_id: str
    booking_date: str
    payee: str
    description: str
    account_number: str
    account_name: str
    debit: Decimal
    credit: Decimal


class JournalBuilder:
    """Creates Excel workbooks with journal entries grouped by currency."""

    def __init__(self, catalog: AccountCatalog):
        self._catalog = catalog

    def build(self, entries: Iterable[LedgerEntry], output_path: Path) -> None:
        grouped: Dict[str, List[JournalRow]] = defaultdict(list)
        for entry in entries:
            amount = entry.amount.copy_abs()
            debit_account = self._catalog.get(entry.debit_account)
            credit_account = self._catalog.get(entry.credit_account)
            debit_name = debit_account.name if debit_account else entry.debit_account
            credit_name = credit_account.name if credit_account else entry.credit_account
            currency = (entry.currency or "UNSPECIFIED").upper()

            grouped[currency].append(
                JournalRow(
                    transaction_id=entry.transaction_id,
                    booking_date=entry.booking_date.isoformat(),
                    payee=entry.payee,
                    description=entry.description,
                    account_number=entry.debit_account,
                    account_name=debit_name,
                    debit=amount,
                    credit=Decimal("0"),
                )
            )
            grouped[currency].append(
                JournalRow(
                    transaction_id=entry.transaction_id,
                    booking_date=entry.booking_date.isoformat(),
                    payee=entry.payee,
                    description=entry.description,
                    account_number=entry.credit_account,
                    account_name=credit_name,
                    debit=Decimal("0"),
                    credit=amount,
                )
            )

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        for currency in sorted(grouped.keys()):
            worksheet = workbook.create_sheet(title=f"{currency} Journal")
            worksheet.append(
                [
                    "Transaction ID",
                    "Date",
                    "Payee",
                    "Description",
                    "Account",
                    "Account name",
                    "Debit",
                    "Credit",
                ]
            )
            for row in grouped[currency]:
                worksheet.append(
                    [
                        row.transaction_id,
                        row.booking_date,
                        row.payee,
                        row.description,
                        row.account_number,
                        row.account_name,
                        float(row.debit),
                        float(row.credit),
                    ]
                )
            self._auto_width(worksheet)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)

    def _auto_width(self, worksheet) -> None:
        for column_cells in worksheet.columns:
            max_length = 0
            column = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = cell.value
                if value is None:
                    continue
                length = len(str(value))
                max_length = max(max_length, length)
            worksheet.column_dimensions[column].width = max(12, min(max_length + 2, 60))
